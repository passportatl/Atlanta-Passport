import io
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

DOMAIN = "fb9954b1-03a2-4563-bd46-f2c651216b72-00-1zry9jh1hlhz4.kirk.replit.dev"
BASE = f"https://{DOMAIN}/stamp"

# (Name, Neighborhood, Category, Passport Offer, db_slug, status)
SPOTS = [
    ("Atlantucky Brewing", "Castleberry Hill", "Drink", "10% off your tab when you show your Atlanta Passport.", "atlantucky", "Live"),
    ("Peachtree Wellness", "Grant Park", "Retail", "10% off for Passport holders.", "peachtree-wellness", "Live"),
    ("Wheelhaus Bikes", "East Atlanta Village", "Rentals", "Rental specials and guided ride options during World Cup season.", "wheelhaus", "Live"),
    ("The Westwood", "West End", "Drink", "A complimentary Bloody Mary when you show your Atlanta Passport.", "westwood", "Live"),
    ("Vickery's Bar & Grill", "Glenwood Park", "Food", "Free dessert with each entrée purchase. Dine-in only. Offer valid through July 31, 2026.", "vickerys", "Live"),
    ("BoxCar at Hop City", "West End", "Retail", "10% off THC Drinks", "boxcar", "Live"),
    ("Hop City at Krog St Market", "Krog", "Retail", "10% off THC Drinks", "hop-city-krog", "Live"),
    ("La Semilla", "Reynoldstown", "Food", "10% off", "la-semilla", "Live"),
    ("Trap Museum", "Westside", "Experiences", "10% off admission when you show your Atlanta Passport.", "trap-museum", "Needs setup — not yet a scannable stamp"),
    ("Varasanos", "Buckhead", "Food", "15% off any food & drink purchase!", "varasanos", "Live"),
    ("Nakato Japanese Restaurant", "Piedmont Heights", "Food", "BOGO: Spicy Tuna Roll", "nakato", "Live"),
]

# (Event, Date, Venue, db_slug)
EVENTS = [
    ("Battle of the Bands", "June 13, 2026", "Atlantucky Brewing", "event-battle-of-the-bands"),
    ("Video Game Prelims + Soccer Tourney", "June 16, 2026", "Atlantucky Brewing", "event-video-game-prelims"),
    ("Hot Sauce Market", "June 20, 2026", "Atlantucky Brewing", "event-hot-sauce-market"),
    ("Post-Match Vibes at Atlantucky", "June 21, 2026", "Atlantucky Brewing", "event-post-match-atlantucky"),
    ("Soccer Video Game Tournament + Wing Eating Comp", "June 22–23, 2026", "Atlantucky Brewing", "event-soccer-gaming-finals"),
]

NAVY = "1A2B4A"
YELLOW = "F4C430"
RED = "C0392B"
CREAM = "FBF3DE"
WHITE = "FFFFFF"

thin = Side(style="thin", color="C9C0A8")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def qr_image(url, px=170):
    api = f"https://api.qrserver.com/v1/create-qr-code/?data={requests.utils.quote(url, safe='')}&size={px}x{px}&margin=8"
    r = requests.get(api, timeout=30)
    r.raise_for_status()
    bio = io.BytesIO(r.content)
    img = XLImage(bio)
    img.width = 110
    img.height = 110
    return img


def style_header(ws, headers, fill_color, fg="FFFFFF"):
    fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color=fg, size=11, name="Calibri")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[ws.max_row].height = 26


def title_row(ws, text, ncols, fill_color):
    ws.append([text])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=ncols)
    c = ws.cell(row=ws.max_row, column=1)
    c.font = Font(bold=True, color="FFFFFF", size=16, name="Calibri")
    c.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[ws.max_row].height = 34


def subtitle_row(ws, text, ncols):
    ws.append([text])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=ncols)
    c = ws.cell(row=ws.max_row, column=1)
    c.font = Font(italic=True, color="555555", size=10, name="Calibri")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[ws.max_row].height = 20


wb = Workbook()

# ---------------- Sheet 1: Stamp QR Codes ----------------
ws1 = wb.active
ws1.title = "Stamp QR Codes"
ncols1 = 7
title_row(ws1, "ATLANTA PASSPORT — Stamp QR Codes", ncols1, NAVY)
subtitle_row(ws1, "Scan a spot's QR code to collect a stamp. 11 partner spots with a Passport offer.", ncols1)
ws1.append([])

headers1 = ["QR Code", "Spot", "Neighborhood", "Category", "Passport Offer", "Scan URL", "Status"]
style_header(ws1, headers1, NAVY)

widths1 = [18, 26, 20, 14, 42, 52, 22]
for i, w in enumerate(widths1, 1):
    ws1.column_dimensions[get_column_letter(i)].width = w

for name, hood, cat, offer, slug, status in SPOTS:
    url = f"{BASE}/{slug}"
    ws1.append(["", name, hood, cat, offer, url, status])
    row = ws1.max_row
    ws1.row_dimensions[row].height = 90
    for col in range(1, ncols1 + 1):
        cell = ws1.cell(row=row, column=col)
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True,
                                   horizontal="left" if col in (5, 6) else "center")
    ws1.cell(row=row, column=2).font = Font(bold=True, size=11)
    sc = ws1.cell(row=row, column=7)
    sc.font = Font(bold=True, color=(RED if status != "Live" else "2E7D32"))
    uc = ws1.cell(row=row, column=6)
    uc.font = Font(color="1A5276", size=9)
    uc.hyperlink = url
    img = qr_image(url)
    ws1.add_image(img, f"A{row}")

# zebra striping
for r in range(5, ws1.max_row + 1):
    if (r - 5) % 2 == 1:
        for col in range(1, ncols1 + 1):
            cell = ws1.cell(row=r, column=col)
            if cell.fill.fill_type != "solid":
                cell.fill = PatternFill(start_color=CREAM, end_color=CREAM, fill_type="solid")

ws1.freeze_panes = "A5"

# ---------------- Sheet 2: Bonus Stamp Events ----------------
ws2 = wb.create_sheet("Bonus Stamp Events")
ncols2 = 5
title_row(ws2, "ATLANTA PASSPORT — Bonus Stamp Events", ncols2, RED)
subtitle_row(ws2, "Scan at the event to add a bonus stamp toward your rewards total. 5 featured events.", ncols2)
ws2.append([])

headers2 = ["QR Code", "Event", "Date", "Venue", "Scan URL"]
style_header(ws2, headers2, RED)

widths2 = [18, 44, 18, 24, 52]
for i, w in enumerate(widths2, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w

for name, date, venue, slug in EVENTS:
    url = f"{BASE}/{slug}"
    ws2.append(["", name, date, venue, url])
    row = ws2.max_row
    ws2.row_dimensions[row].height = 90
    for col in range(1, ncols2 + 1):
        cell = ws2.cell(row=row, column=col)
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True,
                                   horizontal="left" if col == 5 else "center")
    ws2.cell(row=row, column=2).font = Font(bold=True, size=11)
    uc = ws2.cell(row=row, column=5)
    uc.font = Font(color="1A5276", size=9)
    uc.hyperlink = url
    img = qr_image(url)
    ws2.add_image(img, f"A{row}")

for r in range(5, ws2.max_row + 1):
    if (r - 5) % 2 == 1:
        for col in range(1, ncols2 + 1):
            cell = ws2.cell(row=r, column=col)
            if cell.fill.fill_type != "solid":
                cell.fill = PatternFill(start_color=CREAM, end_color=CREAM, fill_type="solid")

ws2.freeze_panes = "A5"

out = "exports/atlanta-passport-qr-codes.xlsx"
wb.save(out)
print("Saved", out)
print("Spots:", len(SPOTS), "Events:", len(EVENTS))
